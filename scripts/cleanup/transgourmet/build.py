import sys, json, csv, re
import os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mapping import TG_ITEMS, TG_ALTERNATIVES
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

oz = json.load(open(sys.argv[1])); okt, mrz = oz['2026-10'], oz['2026-03']
items = {int(r['id']): r for r in csv.DictReader(open(sys.argv[2]), delimiter='\t')}
UNIT = {'LT': 'l', 'KG': 'kg', 'ST': 'stk'}
LE = {'KA': 'Karton', 'BT': 'Beutel', 'FL': 'Flasche', 'PG': 'Packung', 'ST': 'Stück', 'EI': 'Eimer', 'SL': 'Schale',
      'WA': 'Wanne', 'BX': 'Box', 'KG': 'kg', 'DS': 'Dose', 'GL': 'Glas', 'BE': 'Becher', 'SK': 'Sack', 'RL': 'Rolle', 'KI': 'Kiste', 'KT': 'Karton'}

def mode(unit, a):
    if re.match(r'\s*(karton|ktn|ka\b|\d+\s*[x×])', unit, re.I) or not a['anbruch']:
        return 'gebinde' if a['gebinde'] else 'anbruch'
    return 'anbruch'

TO_BASE = {'kg': ('kg', 1), 'g': ('kg', 0.001), 'l': ('l', 1), 'ml': ('l', 0.001)}

def text_size(t):
    """Größe aus der TG-Bezeichnung: '50x125g' → (50, 0.125, 'kg'); '875ml' → (1, 0.875, 'l')."""
    t = t.lower().replace(',', '.')
    m = re.search(r'(\d+)\s*x\s*(\d+(?:\.\d+)?)\s*(kg|g|ml|l)\b', t)
    if m:
        b, f = TO_BASE[m.group(3)]; return int(m.group(1)), float(m.group(2)) * f, b
    ms = re.findall(r'(\d+(?:\.\d+)?)\s*(kg|g|ml|l)\b', t)
    if ms:
        v, u = ms[-1]; b, f = TO_BASE[u]; return 1, float(v) * f, b
    return None

def single(a):
    """Inhalt einer TG-Einzeleinheit als (Menge, Einheit) oder (None, Hinweis)."""
    ts = text_size(a['text1'])
    r, ru = a['rezept'] or 0, UNIT.get(a['rezept_unit'])
    if a['unit'] == 'ST' or (ru == 'stk' and ts and ts[0] > 1):
        return (ts[0] if ts and ts[0] > 1 else 1), 'stk'
    if ru in ('kg', 'l') and r > 0:
        if ts and ts[2] == ru and abs(ts[0] * ts[1] - r) > 0.1 * r:
            return None, f'Inhalt unklar: Bezeichnung {ts[0] * ts[1]:g} {ru}, Ordersatz {r:g} {ru}'
        return r, ru
    if ts:
        return ts[0] * ts[1], ts[2]
    mc = re.search(r'(\d+)\s*(?:st|stk|rl)\b', a['text1'].lower())
    if mc:
        return int(mc.group(1)), 'stk'
    return None, 'Inhalt unbekannt'

def offer(a, m):
    """Bestelleinheit, Preis je Bestelleinheit, Preisbasis und Gebinde."""
    e = a[m] or a['gebinde'] or a['anbruch']
    if a['unit'] == 'KG':                       # nach Gewicht verkauft
        return {'label': 'kg', 'price': round(a['anbruch']['preis_einzel'] if a['anbruch'] else e['preis_einzel'], 2),
                'per': 'kg', 'size': None, 'su': None, 'content': 'Preis je kg', 'note': ''}
    size, su = single(a)
    note = '' if size else su
    if m == 'anbruch':
        label, price, n = LE.get(a['unit'], a['unit']), e['preis_einzel'], 1
    else:
        n = e['inhalt'] or 1
        label = LE.get(a['unit'], a['unit']) if n == 1 else LE.get(e['liefereinheit'], e['liefereinheit'])
        price = e['preis_gebinde']
    content = (f'{n:g}× ' if n != 1 else '') + LE.get(a['unit'], a['unit']) + (f' à {size:g} {su}' if size else '')
    return {'label': label, 'price': round(price, 2), 'per': 'order',
            'size': round(n * size, 4) if size else None, 'su': su if size else None, 'content': content, 'note': note}

def cmp(o):
    if o['per'] == 'kg': return f"{o['price']:.2f} €/kg".replace('.', ',')
    return f"{o['price'] / o['size']:.2f} €/{ {'stk': 'Stück'}.get(o['su'], o['su']) }".replace('.', ',') if o['size'] else ''

def row_for(iid, nr, sich, note, art_alt=False):
    it = items[iid]
    base = [iid, it['name'], it['unit'], it['loc'], it['sup'] or '']
    if not nr:
        return base + ['', '', '', '', '', '', '', '', '', sich, note, 'nein']
    a = okt.get(nr) or mrz.get(nr)
    m = mode(it['unit'], a)
    o = offer(a, m)
    notes = [x for x in [note, o['note']] if x]
    single_unit = not re.match(r'\s*(karton|ktn|ka\b|\d+\s*[x×])', it['unit'], re.I)
    if m == 'gebinde' and single_unit and (a['gebinde']['inhalt'] or 1) > 1:
        notes.append(f"TG liefert nur ganze {o['label']} ({o['content']}) – passt zur Einheit „{it['unit']}“?")
        sich = 'prüfen'
    pm = ''
    if nr in mrz:
        a2 = mrz[nr]; pm = offer(a2, m if a2[m] else ('gebinde' if a2['gebinde'] else 'anbruch'))['price']
    if o['note'] and sich == 'sicher': sich = 'prüfen'
    return base + [nr, a['text1'], o['label'], o['content'], f"{o['size']:g} {o['su']}" if o['size'] else '',
                   o['price'], o['per'], cmp(o), pm, sich, ' · '.join(notes), 'ja' if (sich == 'sicher' and not art_alt) else '']

def main():
    wb = openpyxl.Workbook()
    H1 = ['ID', 'Name in Orderlauf', 'Einheit', 'Lager', 'Lieferanten (Prio)', 'TG Art.-Nr.', 'TG-Bezeichnung (→ Bestellbezeichnung)',
          'Bestelleinheit TG', 'Inhalt', 'Gebinde (für €/kg)', 'Preis Okt. 2026 € (netto)', 'Preis je', 'Vergleichspreis', 'Preis März 2026 €',
          'Sicherheit', 'Hinweis', 'Übernehmen? (ja/nein)']
    def sheet(ws, title, rows, help_text):
        ws.title = title
        ws.append([help_text]); ws['A1'].font = Font(italic=True); ws.append([])
        ws.append(H1)
        for c in ws[3]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBE4'); c.alignment = Alignment(wrap_text=True, vertical='top')
        for r in rows: ws.append(r)
        for r in ws.iter_rows(min_row=4):
            r[15].alignment = Alignment(wrap_text=True, vertical='top')
            s = r[14].value
            r[14].fill = PatternFill('solid', fgColor={'sicher': 'E3F2E1', 'prüfen': 'FFF4D6', 'kein Treffer': 'F6DADA'}.get(s, 'FFF4D6'))
            r[16].fill = PatternFill('solid', fgColor='FFFBEA')
        dv = DataValidation(type='list', formula1='"ja,nein"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'Q4:Q{ws.max_row}')
        for col, w in zip('ABCDEFGHIJKLMNOPQ', [5, 32, 16, 6, 20, 10, 32, 11, 22, 12, 10, 8, 14, 10, 11, 60, 12]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'C4'; ws.auto_filter.ref = f'A3:Q{ws.max_row}'

    sheet(wb.active, 'Transgourmet-Artikel', [row_for(*t) for t in TG_ITEMS],
          'Artikel, die heute (auch) bei Transgourmet bestellt werden. „ja“ = Bestellbezeichnung, Art.-Nr., Gebinde und beide Preise (Listenpreis) übernehmen. Namen in Orderlauf bleiben unverändert.')
    sheet(wb.create_sheet(), 'Mögliche Alternativen', [row_for(iid, nr, 'prüfen', note, True) for iid, nr, note in TG_ALTERNATIVES],
          'Artikel, die du heute woanders bestellst, gibt es ähnlich bei Transgourmet. „ja“ = Transgourmet als Alternativ-Lieferant mit niedrigerer Priorität ergänzen (inkl. Bestelldaten und Preis). Austauschbarkeit bitte selbst beurteilen.')
    wb.save(sys.argv[3])
    print('ok', sys.argv[3])


if __name__ == '__main__':
    main()
